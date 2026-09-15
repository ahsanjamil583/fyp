"""Import historical order sheets into the unified transaction model.

The flow is deliberately two-step. ``preview_order_sheet`` parses the upload, normalizes
headers, groups the rows into orders and reports what it found and what it could not
read, without writing anything. The owner then confirms, and ``confirm_order_import``
writes the rows they accepted. Nothing is written on the strength of a file upload alone,
which is what makes a mis-mapped column recoverable.

Imported orders are real transactions with ``source: "imported"``: they carry their
original date and number, they show up in the owner's queue, in analytics and in customer
history. They deliberately do **not** move stock - the goods left the shelf whenever the
original sale happened, often months ago, and replaying that now would corrupt today's
inventory.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any

from bson import ObjectId
from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook

from app.core.object_ids import parse_object_id, serialize_document
from app.core.permissions import get_owned_tenant_or_403
from app.db.mongodb import get_database
from app.services.customer_service import sync_customer_stats_for_transaction
from app.services.localization_service import normalize_optional_email, normalize_optional_pk_phone_or_blank
from app.services.transaction_number_service import generate_transaction_number
from app.services.transaction_workflow_service import (
    get_allowed_payment_statuses,
    get_allowed_statuses,
)

IMPORTED_SOURCE = "imported"
MAX_IMPORT_ROWS = 5000
MAX_IMPORT_FILE_BYTES = 8 * 1024 * 1024

SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".csv")

# What the owner is told to put in the sheet, and every spelling this parser accepts.
IMPORT_COLUMNS: list[dict[str, Any]] = [
    {"key": "orderNumber", "label": "Order number", "required": False, "aliases": ["order number", "order no", "order id", "invoice", "invoice number", "invoice no", "receipt number", "bill number"]},
    {"key": "orderDate", "label": "Order date", "required": False, "aliases": ["order date", "date", "created at", "created on", "invoice date", "bill date", "sale date"]},
    {"key": "customerName", "label": "Customer name", "required": False, "aliases": ["customer name", "customer", "client", "client name", "buyer", "name"]},
    {"key": "customerPhone", "label": "Customer phone", "required": False, "aliases": ["customer phone", "phone", "mobile", "contact", "contact number", "phone number"]},
    {"key": "customerEmail", "label": "Customer email", "required": False, "aliases": ["customer email", "email", "email address"]},
    {"key": "itemName", "label": "Item name", "required": False, "aliases": ["item name", "item", "product", "product name", "description", "particulars", "service"]},
    {"key": "sku", "label": "SKU", "required": False, "aliases": ["sku", "item code", "product code", "code"]},
    {"key": "quantity", "label": "Quantity", "required": False, "aliases": ["quantity", "qty", "units", "pcs"]},
    {"key": "unitPrice", "label": "Unit price", "required": False, "aliases": ["unit price", "price", "rate", "unit rate", "item price"]},
    {"key": "discount", "label": "Discount", "required": False, "aliases": ["discount", "discount amount", "less"]},
    {"key": "tax", "label": "Tax", "required": False, "aliases": ["tax", "tax amount", "gst", "sales tax", "vat"]},
    {"key": "totalAmount", "label": "Total amount", "required": False, "aliases": ["total amount", "total", "grand total", "amount", "net amount", "payable"]},
    {"key": "paymentMethod", "label": "Payment method", "required": False, "aliases": ["payment method", "payment", "paid via", "mode of payment", "payment mode"]},
    {"key": "paymentStatus", "label": "Payment status", "required": False, "aliases": ["payment status", "paid status", "payment state"]},
    {"key": "orderStatus", "label": "Order status", "required": False, "aliases": ["order status", "status", "state"]},
    {"key": "notes", "label": "Notes", "required": False, "aliases": ["notes", "note", "remarks", "comment", "comments"]},
]

_ALIAS_TO_KEY: dict[str, str] = {}
for column in IMPORT_COLUMNS:
    _ALIAS_TO_KEY[column["label"].lower()] = column["key"]
    _ALIAS_TO_KEY[column["key"].lower()] = column["key"]
    for alias in column["aliases"]:
        _ALIAS_TO_KEY[alias] = column["key"]

PAYMENT_STATUS_MAP = {
    "paid": "paid",
    "complete": "paid",
    "completed": "paid",
    "settled": "paid",
    "received": "paid",
    "unpaid": "unpaid",
    "due": "unpaid",
    "pending": "unpaid",
    "credit": "unpaid",
    "partial": "partially_paid",
    "partially paid": "partially_paid",
    "partially_paid": "partially_paid",
    "cod": "cod",
    "cash on delivery": "cod",
    "refund": "refunded",
    "refunded": "refunded",
    "cancelled": "unpaid",
}

ORDER_STATUS_MAP = {
    "completed": "completed",
    "complete": "completed",
    "delivered": "completed",
    "done": "completed",
    "closed": "completed",
    "fulfilled": "completed",
    "pending": "pending",
    "new": "pending",
    "open": "pending",
    "confirmed": "confirmed",
    "processing": "processing",
    "in progress": "processing",
    "ready": "ready",
    "cancelled": "cancelled",
    "canceled": "cancelled",
    "returned": "cancelled",
}

_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%d-%b-%Y",
    "%d %b %Y",
    "%d %B %Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%d/%m/%Y %H:%M",
    "%m/%d/%Y %H:%M",
)


def import_column_guide() -> list[dict[str, Any]]:
    """What the upload screen shows before a file is chosen."""
    return [
        {"key": column["key"], "label": column["label"], "required": column["required"], "examples": column["aliases"][:4]}
        for column in IMPORT_COLUMNS
    ]


def normalize_import_header(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    spaced = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", text)
    cleaned = re.sub(r"[^a-zA-Z0-9]+", " ", spaced).strip().lower()
    return _ALIAS_TO_KEY.get(cleaned, "")


def _text(value: Any, limit: int = 200) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()[:limit]


def _number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    if cleaned in {"", "-", ".", "-."}:
        return default
    try:
        return float(cleaned)
    except ValueError:
        return default


def parse_import_date(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def map_payment_status(value: Any) -> str:
    normalized = re.sub(r"[^a-z ]", " ", str(value or "").strip().lower())
    normalized = re.sub(r"\s+", " ", normalized).strip()
    if not normalized:
        return "unpaid"
    return PAYMENT_STATUS_MAP.get(normalized, PAYMENT_STATUS_MAP.get(normalized.replace(" ", "_"), "unpaid"))


def map_order_status(value: Any) -> str:
    normalized = re.sub(r"[^a-z ]", " ", str(value or "").strip().lower())
    normalized = re.sub(r"\s+", " ", normalized).strip()
    if not normalized:
        return "completed"
    return ORDER_STATUS_MAP.get(normalized, "completed")


async def _read_sheet_rows(file: UploadFile) -> tuple[list[str], list[list[Any]]]:
    filename = (file.filename or "").lower()
    if not filename.endswith(SUPPORTED_EXTENSIONS):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Upload a .xlsx, .xlsm or .csv file.")

    payload = await file.read(MAX_IMPORT_FILE_BYTES + 1)
    if len(payload) > MAX_IMPORT_FILE_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Order sheets must be 8 MB or smaller.")
    if not payload:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="That file is empty.")

    if filename.endswith(".csv"):
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = payload.decode("latin-1")
        reader = csv.reader(io.StringIO(text))
        rows = [row for row in reader]
    elif filename.endswith(".xls"):
        # openpyxl reads the OOXML format only. Saying so is more useful than a parser
        # traceback, because "save as .xlsx" is a ten-second fix in Excel.
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="The old .xls format is not supported. Open the file in Excel and save it as .xlsx or .csv, then upload again.",
        )
    else:
        try:
            workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
        except Exception as exc:  # openpyxl raises a wide range of parse errors
            raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="That spreadsheet could not be read. Save it again as .xlsx and retry.") from exc
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    if not rows:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="That file has no rows.")
    if len(rows) - 1 > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"Split the file: at most {MAX_IMPORT_ROWS} data rows can be imported at once.")

    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    return headers, rows[1:]


def _row_key(parsed: dict) -> str:
    """Group multi-line orders.

    Sheets export one line per item, repeating the order number. Rows that carry an order
    number group by it; rows without one become an order of their own, because guessing
    that two numberless rows belong together would silently merge unrelated sales.
    """
    order_number = parsed["orderNumber"].strip().lower()
    if order_number:
        return f"num:{order_number}"
    return f"row:{parsed['rowNumber']}"


def _parse_rows(headers: list[str], data_rows: list[list[Any]]) -> tuple[list[dict], list[dict], dict[str, str]]:
    header_map: dict[str, int] = {}
    mapping: dict[str, str] = {}
    for index, header in enumerate(headers):
        key = normalize_import_header(header)
        if key and key not in header_map:
            header_map[key] = index
            mapping[header or f"Column {index + 1}"] = key

    errors: list[dict] = []
    grouped: dict[str, dict] = {}

    if not header_map:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No recognizable column headers were found. The first row must name the columns, for example: Order number, Order date, Customer name, Item name, Quantity, Unit price, Total amount.",
        )

    for offset, row in enumerate(data_rows):
        row_number = offset + 2  # 1-based, and row 1 is the header
        def cell(key: str) -> Any:
            index = header_map.get(key)
            return row[index] if index is not None and index < len(row) else None

        parsed = {
            "rowNumber": row_number,
            "orderNumber": _text(cell("orderNumber"), 80),
            "orderDateRaw": _text(cell("orderDate"), 60),
            "customerName": _text(cell("customerName"), 120),
            "customerPhone": _text(cell("customerPhone"), 30),
            "customerEmail": _text(cell("customerEmail"), 200),
            "itemName": _text(cell("itemName"), 200),
            "sku": _text(cell("sku"), 80),
            "quantity": _number(cell("quantity"), 1),
            "unitPrice": _number(cell("unitPrice"), 0),
            "discount": _number(cell("discount"), 0),
            "tax": _number(cell("tax"), 0),
            "totalAmount": _number(cell("totalAmount"), 0),
            "paymentMethod": _text(cell("paymentMethod"), 40),
            "paymentStatus": _text(cell("paymentStatus"), 40),
            "orderStatus": _text(cell("orderStatus"), 40),
            "notes": _text(cell("notes"), 1000),
        }

        order_date = parse_import_date(cell("orderDate"))
        if parsed["orderDateRaw"] and order_date is None:
            errors.append({"row": row_number, "field": "orderDate", "message": f"Could not read the date '{parsed['orderDateRaw']}'. Use a format like 2025-03-18 or 18/03/2025."})
            continue

        line_total = parsed["quantity"] * parsed["unitPrice"]
        if not parsed["itemName"] and line_total <= 0 and parsed["totalAmount"] <= 0:
            errors.append({"row": row_number, "field": "itemName", "message": "This row has no item and no amount, so there is nothing to import."})
            continue
        if parsed["quantity"] < 0 or parsed["unitPrice"] < 0 or parsed["totalAmount"] < 0:
            errors.append({"row": row_number, "field": "amount", "message": "Quantities and amounts cannot be negative."})
            continue

        key = _row_key(parsed)
        group = grouped.get(key)
        if not group:
            group = {
                "rowNumbers": [],
                "rowNumber": row_number,
                "orderNumber": parsed["orderNumber"],
                "orderDate": order_date.isoformat() if order_date else "",
                "customerName": parsed["customerName"],
                "customerPhone": parsed["customerPhone"],
                "customerEmail": parsed["customerEmail"],
                "items": [],
                "discount": 0.0,
                "tax": 0.0,
                "totalAmount": 0.0,
                "sheetTotal": 0.0,
                "paymentMethod": parsed["paymentMethod"],
                "paymentStatus": parsed["paymentStatus"],
                "orderStatus": parsed["orderStatus"],
                "notes": parsed["notes"],
            }
            grouped[key] = group

        group["rowNumbers"].append(row_number)
        group["customerName"] = group["customerName"] or parsed["customerName"]
        group["customerPhone"] = group["customerPhone"] or parsed["customerPhone"]
        group["customerEmail"] = group["customerEmail"] or parsed["customerEmail"]
        group["paymentMethod"] = group["paymentMethod"] or parsed["paymentMethod"]
        group["paymentStatus"] = group["paymentStatus"] or parsed["paymentStatus"]
        group["orderStatus"] = group["orderStatus"] or parsed["orderStatus"]
        if parsed["notes"] and parsed["notes"] not in group["notes"]:
            group["notes"] = (group["notes"] + " " + parsed["notes"]).strip()[:1000]
        if not group["orderDate"] and order_date:
            group["orderDate"] = order_date.isoformat()

        unit_price = parsed["unitPrice"]
        quantity = parsed["quantity"] or 1
        if unit_price <= 0 and parsed["totalAmount"] > 0 and quantity:
            # Sheets that only carry a line total still describe a sale; back out the rate.
            unit_price = round(parsed["totalAmount"] / quantity, 2)
        group["items"].append(
            {
                "name": parsed["itemName"] or "Imported item",
                "sku": parsed["sku"],
                "quantity": quantity,
                "unitPrice": unit_price,
            }
        )
        group["discount"] += parsed["discount"]
        group["tax"] += parsed["tax"]
        group["sheetTotal"] += parsed["totalAmount"]

    rows_out = []
    for group in grouped.values():
        computed = sum(item["quantity"] * item["unitPrice"] for item in group["items"])
        subtotal = round(computed, 2)
        discount = round(group["discount"], 2)
        tax = round(group["tax"], 2)
        derived_total = round(max(0.0, subtotal - discount + tax), 2)
        # A stated total wins over the arithmetic: the sheet is the record of what was
        # actually charged, rounding and all.
        total = round(group["sheetTotal"], 2) if group["sheetTotal"] > 0 else derived_total
        group["subtotal"] = subtotal
        group["discount"] = discount
        group["tax"] = tax
        group["totalAmount"] = total
        group["totalMismatch"] = bool(group["sheetTotal"] > 0 and abs(derived_total - total) > 0.5)
        group.pop("sheetTotal", None)
        rows_out.append(group)

    rows_out.sort(key=lambda row: row["rowNumber"])
    return rows_out, errors, mapping


def _duplicate_signature(tenant_oid: ObjectId, row: dict) -> dict:
    """Match on number when there is one, otherwise on date + customer + amount."""
    order_number = (row.get("orderNumber") or "").strip()
    if order_number:
        return {"tenantId": tenant_oid, "importSignature.orderNumber": order_number.lower()}
    return {
        "tenantId": tenant_oid,
        "importSignature.fingerprint": _fingerprint(row),
    }


def _fingerprint(row: dict) -> str:
    date_part = (row.get("orderDate") or "")[:10]
    name_part = re.sub(r"\s+", " ", (row.get("customerName") or "").strip().lower())
    phone_part = re.sub(r"\D", "", row.get("customerPhone") or "")
    total_part = f"{float(row.get('totalAmount') or 0):.2f}"
    return f"{date_part}|{name_part}|{phone_part}|{total_part}"


async def _flag_duplicates(tenant_oid: ObjectId, rows: list[dict]) -> int:
    db = get_database()
    duplicate_count = 0
    seen_in_file: set[str] = set()
    for row in rows:
        signature = (row.get("orderNumber") or "").strip().lower() or _fingerprint(row)
        if signature in seen_in_file:
            row["isDuplicate"] = True
            row["duplicateReason"] = "The same order appears earlier in this file."
            duplicate_count += 1
            continue
        seen_in_file.add(signature)
        existing = await db.transactions.find_one(_duplicate_signature(tenant_oid, row), {"_id": 1, "transactionNumber": 1})
        if existing:
            row["isDuplicate"] = True
            row["duplicateReason"] = f"Already imported as {existing.get('transactionNumber', 'an existing order')}."
            duplicate_count += 1
        else:
            row["isDuplicate"] = False
            row["duplicateReason"] = ""
    return duplicate_count


async def _ensure_import_access(tenant_id: str, user: dict) -> tuple[ObjectId, dict]:
    tenant_oid = parse_object_id(tenant_id, "tenantId")
    tenant = await get_owned_tenant_or_403(tenant_oid, user)
    return tenant_oid, tenant


async def preview_order_sheet(tenant_id: str, file: UploadFile, user: dict) -> dict:
    tenant_oid, _ = await _ensure_import_access(tenant_id, user)
    headers, data_rows = await _read_sheet_rows(file)
    rows, errors, mapping = _parse_rows(headers, data_rows)
    duplicates = await _flag_duplicates(tenant_oid, rows)

    return {
        "fileName": file.filename or "orders.csv",
        "detectedColumns": [{"header": header, "mappedTo": key} for header, key in mapping.items()],
        "unmappedColumns": [header for header in headers if header and not normalize_import_header(header)],
        "totalSheetRows": len(data_rows),
        "orders": rows,
        "errors": errors,
        "summary": {
            "readyRows": sum(1 for row in rows if not row.get("isDuplicate")),
            "duplicateRows": duplicates,
            "failedRows": len(errors),
            "totalValue": round(sum(float(row.get("totalAmount") or 0) for row in rows if not row.get("isDuplicate")), 2),
        },
    }


async def confirm_order_import(tenant_id: str, payload, user: dict) -> dict:
    db = get_database()
    tenant_oid, tenant = await _ensure_import_access(tenant_id, user)
    if not payload.rows:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="There are no rows to import.")
    if len(payload.rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"At most {MAX_IMPORT_ROWS} orders can be imported at once.")

    now = datetime.now(timezone.utc)
    import_record = {
        "tenantId": tenant_oid,
        "fileName": str(payload.fileName or "imported-orders")[:255],
        "totalRows": len(payload.rows),
        "successCount": 0,
        "skippedCount": 0,
        "errorCount": 0,
        "errors": [],
        "importedValue": 0.0,
        "status": "running",
        "createdBy": user["_id"],
        "createdByName": user.get("fullName", ""),
        "createdAt": now,
        "updatedAt": now,
    }
    import_record["_id"] = (await db.order_imports.insert_one(import_record)).inserted_id

    created: list[dict] = []
    skipped: list[dict] = []
    errors: list[dict] = []
    imported_value = 0.0

    for row in payload.rows:
        row_data = row.model_dump()
        row_number = int(row_data.get("rowNumber") or 0)
        try:
            if payload.skipDuplicates and await db.transactions.find_one(_duplicate_signature(tenant_oid, row_data), {"_id": 1}):
                skipped.append({"row": row_number, "orderNumber": row_data.get("orderNumber", ""), "message": "Already present, skipped."})
                continue

            transaction = await _build_imported_transaction(tenant_oid, tenant, row_data, user, import_record["_id"], now)
            created.append(transaction)
            imported_value += float((transaction.get("pricing") or {}).get("total") or 0)
        except HTTPException as exc:
            errors.append({"row": row_number, "orderNumber": row_data.get("orderNumber", ""), "message": exc.detail if isinstance(exc.detail, str) else "Row rejected."})
        except Exception as exc:
            errors.append({"row": row_number, "orderNumber": row_data.get("orderNumber", ""), "message": str(exc)[:300]})

    final_status = "completed" if not errors else ("failed" if not created else "completed_with_errors")
    await db.order_imports.update_one(
        {"_id": import_record["_id"]},
        {
            "$set": {
                "successCount": len(created),
                "skippedCount": len(skipped),
                "errorCount": len(errors),
                "errors": errors[:500],
                "skipped": skipped[:500],
                "importedValue": round(imported_value, 2),
                "status": final_status,
                "completedAt": datetime.now(timezone.utc),
                "updatedAt": datetime.now(timezone.utc),
            }
        },
    )
    saved = await db.order_imports.find_one({"_id": import_record["_id"]})
    return {
        "import": serialize_document(saved),
        "createdOrders": [
            {
                "transactionNumber": transaction.get("transactionNumber", ""),
                "total": (transaction.get("pricing") or {}).get("total", 0),
                "customerName": (transaction.get("customerSnapshot") or {}).get("name", ""),
                "createdAt": transaction.get("createdAt").isoformat() if hasattr(transaction.get("createdAt"), "isoformat") else transaction.get("createdAt"),
            }
            for transaction in created[:200]
        ],
    }


async def _build_imported_transaction(tenant_oid: ObjectId, tenant: dict, row: dict, user: dict, import_id: ObjectId, run_at: datetime) -> dict:
    db = get_database()
    items = row.get("items") or []
    if not items:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="This order has no item lines.")

    lines = []
    subtotal = 0.0
    for entry in items:
        quantity = max(float(entry.get("quantity") or 1), 0.0) or 1.0
        unit_price = max(float(entry.get("unitPrice") or 0), 0.0)
        line_subtotal = round(quantity * unit_price, 2)
        subtotal += line_subtotal
        lines.append(
            {
                "itemId": None,
                "name": str(entry.get("name") or "Imported item")[:200],
                "sku": str(entry.get("sku") or "")[:80],
                "quantity": quantity,
                "unitPrice": unit_price,
                "currency": ((tenant.get("settings") or {}).get("currency")) or "PKR",
                "subtotal": line_subtotal,
                "isCustomItem": True,
                # Historical lines never move stock; see the module docstring.
                "stockSnapshot": {"tracked": False, "scope": "not_tracked", "available": True, "message": "Imported from a historical sheet."},
            }
        )

    subtotal = round(subtotal, 2)
    discount = round(max(float(row.get("discount") or 0), 0.0), 2)
    tax = round(max(float(row.get("tax") or 0), 0.0), 2)
    total = round(float(row.get("totalAmount") or 0), 2) or round(max(0.0, subtotal - discount + tax), 2)

    order_status = map_order_status(row.get("orderStatus"))
    if order_status not in get_allowed_statuses("order"):
        order_status = "completed"
    payment_status = map_payment_status(row.get("paymentStatus"))
    if payment_status not in get_allowed_payment_statuses("order"):
        payment_status = "unpaid"

    order_date = parse_import_date(row.get("orderDate")) or run_at
    paid_amount = total if payment_status in {"paid", "cod"} else 0.0

    transaction = {
        "tenantId": tenant_oid,
        "branchId": None,
        "customerId": None,
        "customerUserId": None,
        "customerProfileId": None,
        "transactionType": "order",
        "transactionNumber": str(row.get("orderNumber") or "").strip()[:80] or await generate_transaction_number(tenant_oid, "order"),
        "source": IMPORTED_SOURCE,
        "status": order_status,
        "items": lines,
        "pricing": {
            "subtotal": subtotal,
            "discount": discount,
            "tax": tax,
            "serviceCharge": 0,
            "deliveryFee": 0,
            "total": total,
        },
        "paymentStatus": payment_status,
        "paymentSummary": {
            "total": total,
            "paid": paid_amount,
            "cod": 0,
            "pending": 0,
            "rejected": 0,
            "refunded": 0,
            "balance": round(max(0.0, total - paid_amount), 2),
        },
        "payment": {
            "method": str(row.get("paymentMethod") or "")[:40],
            "methodLabel": str(row.get("paymentMethod") or "Not recorded")[:40],
        },
        "fulfillment": {"type": "none", "address": {}, "serviceType": "imported"},
        "customerSnapshot": {
            "name": str(row.get("customerName") or "").strip()[:120] or "Imported customer",
            "phone": normalize_optional_pk_phone_or_blank(row.get("customerPhone") or ""),
            "email": normalize_optional_email(row.get("customerEmail") or ""),
        },
        "notes": str(row.get("notes") or "")[:1000],
        "internalNotes": "",
        "customFields": {},
        # Stock was moved when the original sale happened, so this row has nothing to
        # reserve or deduct. Marking it explicitly keeps it out of reconciliation reports.
        "inventoryStatus": "not_required",
        "importSignature": {
            "orderNumber": str(row.get("orderNumber") or "").strip().lower(),
            "fingerprint": _fingerprint(row),
            "importId": import_id,
            "sheetRow": int(row.get("rowNumber") or 0),
        },
        "statusHistory": [
            {
                "field": "status",
                "from": None,
                "to": order_status,
                "note": "Imported from a historical order sheet.",
                "changedAt": run_at,
                "changedByUserId": user["_id"],
            }
        ],
        "createdBy": user["_id"],
        "createdAt": order_date,
        "importedAt": run_at,
        "updatedAt": run_at,
    }

    try:
        transaction["_id"] = (await db.transactions.insert_one(transaction)).inserted_id
    except Exception:
        # The tenant/transactionNumber index is unique, so a sheet that reuses a number
        # the business already issued gets a suffixed one rather than failing the import.
        transaction["transactionNumber"] = f"{transaction['transactionNumber']}-IMP{int(run_at.timestamp()) % 100000}"
        transaction["_id"] = (await db.transactions.insert_one(transaction)).inserted_id

    if transaction["customerSnapshot"]["phone"] or transaction["customerSnapshot"]["email"]:
        await sync_customer_stats_for_transaction(tenant, transaction)
    return transaction


async def list_order_imports(tenant_id: str, user: dict, page: int = 1, limit: int = 20) -> dict:
    db = get_database()
    tenant_oid, _ = await _ensure_import_access(tenant_id, user)
    page = max(page, 1)
    limit = min(max(limit, 1), 50)
    query = {"tenantId": tenant_oid}
    total = await db.order_imports.count_documents(query)
    cursor = db.order_imports.find(query).sort("createdAt", -1).skip((page - 1) * limit).limit(limit)
    items = []
    async for row in cursor:
        record = serialize_document(row) or {}
        record["errors"] = (record.get("errors") or [])[:25]
        record["skipped"] = (record.get("skipped") or [])[:25]
        items.append(record)
    return {
        "items": items,
        "pagination": {"page": page, "limit": limit, "total": total, "totalPages": (total + limit - 1) // limit},
    }


async def get_order_import_error_report(tenant_id: str, import_id: str, user: dict) -> tuple[str, str]:
    """Return (filename, csv text) of the rows an import could not take."""
    db = get_database()
    tenant_oid, _ = await _ensure_import_access(tenant_id, user)
    record = await db.order_imports.find_one({"_id": parse_object_id(import_id, "importId"), "tenantId": tenant_oid})
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Import run not found.")

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Row", "Order number", "Outcome", "Message"])
    for error in record.get("errors") or []:
        writer.writerow([error.get("row", ""), error.get("orderNumber", ""), "failed", error.get("message", "")])
    for skipped in record.get("skipped") or []:
        writer.writerow([skipped.get("row", ""), skipped.get("orderNumber", ""), "skipped", skipped.get("message", "")])

    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "-", str(record.get("fileName") or "import")).strip("-") or "import"
    return f"{safe_name}-errors.csv", buffer.getvalue()
